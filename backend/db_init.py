# -*- coding: utf-8 -*-
"""
Planogram SQLite Veritabanı İlklendirme ve Excel Taşıma Betiği.
"""
import os
import sys
import openpyxl

from database import connect_db

BACKEND_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BACKEND_DIR, "data")
DB_PATH = os.path.join(DATA_DIR, "planogram.db")
TEMP_DB_PATH = os.path.join(DATA_DIR, ".planogram.db.tmp")

QUARTER_MAP = {
    'OCAK': 'Q1', 'ŞUBAT': 'Q1', 'MART': 'Q1',
    'NİSAN': 'Q2', 'MAYIS': 'Q2', 'HAZİRAN': 'Q2', 'HAZİRAN ': 'Q2',
    'TEMMUZ': 'Q3', 'AĞUSTOS': 'Q3', 'EYLÜL': 'Q3',
    'EKİM': 'Q4', 'KASIM': 'Q4', 'ARALIK': 'Q4',
}

STATION_PASSWORD = "shell2025"
ADMIN_USERNAME = "admin"
ADMIN_PASSWORD = "admin123"
_HASH_METHOD = "pbkdf2:sha256"

DEFAULT_CHOCOLATE_WEIGHTS = {
    "satis": 0.22,
    "birim_kar": 0.15,
    "c2": 0.10,
    "ciro_pay": 0.15,
    "niel_gas": 0.20,
    "niel_spm": 0.08,
    "deli2go": 0.10,
}

# Proje kılavuzunda 3 Modül Raf 5'e ayrılan 17 paket ürün.
PACKAGE_PRODUCT_IDS = {
    "8843355507", "7814055142", "7967767083", "8348097760",
    "6139774956", "6654999207", "5795107404", "7397429079",
    "7699273039", "7638440038", "7403928141", "8212920593",
    "9093651530", "8865815672", "9449350786", "9698924284",
    "9866080176",
}


def create_tables(conn):
    cursor = conn.cursor()

    # 1. İstasyonlar tablosu
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS stations (
        roc INTEGER PRIMARY KEY,
        name TEXT NOT NULL
    )
    """)

    # 2. Gondol ürünleri tablosu
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS gondol_products (
        product_name TEXT PRIMARY KEY,
        cluster TEXT NOT NULL,
        default_shelf INTEGER,
        width_cm REAL
    )
    """)

    # 3. Gondol satış verileri tablosu
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS sales (
        roc INTEGER NOT NULL,
        quarter TEXT NOT NULL CHECK (quarter IN ('Q1', 'Q2', 'Q3', 'Q4')),
        product_name TEXT NOT NULL,
        qty REAL NOT NULL DEFAULT 0,
        PRIMARY KEY (roc, quarter, product_name),
        FOREIGN KEY (roc) REFERENCES stations(roc) ON UPDATE CASCADE ON DELETE CASCADE,
        FOREIGN KEY (product_name) REFERENCES gondol_products(product_name) ON UPDATE CASCADE ON DELETE CASCADE
    )
    """)

    # 4. Çikolata ürünleri tablosu
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS chocolate_products (
        urun_id TEXT PRIMARY KEY,
        sku TEXT NOT NULL,
        short_sku TEXT,
        brand TEXT,
        sub_category TEXT,
        width_cm REAL NOT NULL CHECK (width_cm > 0),
        score REAL,
        yerlesim TEXT,
        urun_tipi TEXT NOT NULL DEFAULT 'Mevcut',
        min_facing INTEGER NOT NULL DEFAULT 1 CHECK (min_facing >= 1),
        max_facing INTEGER NOT NULL DEFAULT 1 CHECK (max_facing >= min_facing),
        current_facing INTEGER NOT NULL DEFAULT 1 CHECK (current_facing >= 0),
        is_package INTEGER NOT NULL DEFAULT 0 CHECK (is_package IN (0, 1))
    )
    """)

    # 5. Çikolata metrikleri tablosu
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS chocolate_metrics (
        urun_id TEXT PRIMARY KEY,
        satis REAL,
        birim_kar REAL,
        c2 REAL,
        ciro_pay REAL,
        niel_gas REAL,
        niel_spm REAL,
        deli2go REAL,
        FOREIGN KEY (urun_id) REFERENCES chocolate_products(urun_id) ON UPDATE CASCADE ON DELETE CASCADE
    )
    """)

    # 6. İstasyon bazlı çikolata satışları tablosu
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS chocolate_station_sales (
        roc INTEGER NOT NULL,
        urun_id TEXT NOT NULL,
        q1 REAL NOT NULL DEFAULT 0,
        q2 REAL NOT NULL DEFAULT 0,
        q3 REAL NOT NULL DEFAULT 0,
        q4 REAL NOT NULL DEFAULT 0,
        PRIMARY KEY (roc, urun_id),
        FOREIGN KEY (roc) REFERENCES stations(roc) ON UPDATE CASCADE ON DELETE CASCADE,
        FOREIGN KEY (urun_id) REFERENCES chocolate_products(urun_id) ON UPDATE CASCADE ON DELETE CASCADE
    )
    """)

    # 7. Çikolata modül ve raf konfigürasyonu
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS chocolate_shelves (
        module TEXT NOT NULL CHECK (module IN ('2', '3')),
        shelf_no INTEGER NOT NULL CHECK (shelf_no BETWEEN 1 AND 5),
        width_cm REAL NOT NULL CHECK (width_cm > 0),
        multiplier REAL NOT NULL CHECK (multiplier > 0),
        priority_rank INTEGER NOT NULL CHECK (priority_rank BETWEEN 1 AND 5),
        package_only INTEGER NOT NULL DEFAULT 0 CHECK (package_only IN (0, 1)),
        PRIMARY KEY (module, shelf_no),
        UNIQUE (module, priority_rank)
    )
    """)

    # 8. Varsayılan skor profili
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS chocolate_scoring_weights (
        metric TEXT PRIMARY KEY,
        weight REAL NOT NULL CHECK (weight >= 0)
    )
    """)

    # 9. Kullanıcının arayüzden eklediği ürünler; kaynak Excel tablolarını değiştirmez.
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS custom_products (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        target TEXT NOT NULL CHECK (target IN ('gondol', 'choco3', 'choco2')),
        sku TEXT NOT NULL,
        cluster_key TEXT,
        display_name TEXT,
        color TEXT,
        shelf TEXT,
        width_cm REAL NOT NULL DEFAULT 0,
        sales REAL,
        marka TEXT,
        alt_kategori TEXT,
        tahmini_skor REAL,
        yerlesim TEXT,
        urun_tipi TEXT,
        created_by TEXT,
        created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
    )
    """)

    # 9. Kullanıcılar tablosu (Şifreler veritabanında saklanmayacaktır)
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS users (
        username TEXT PRIMARY KEY,
        role TEXT NOT NULL CHECK (role IN ('admin', 'station')),
        display_name TEXT NOT NULL,
        roc INTEGER,
        CHECK ((role = 'admin' AND roc IS NULL) OR (role = 'station' AND roc IS NOT NULL)),
        FOREIGN KEY (roc) REFERENCES stations(roc) ON UPDATE CASCADE ON DELETE CASCADE
    )
    """)

    cursor.execute("CREATE INDEX IF NOT EXISTS idx_sales_roc_quarter ON sales (roc, quarter)")
    cursor.execute(
        "CREATE INDEX IF NOT EXISTS idx_chocolate_station_sales_roc "
        "ON chocolate_station_sales (roc)"
    )

    shelf_rows = []
    multipliers = {1: 1.30, 2: 1.10, 3: 0.95, 4: 0.80, 5: 0.65}
    priority_ranks = {1: 1, 3: 2, 2: 3, 4: 4, 5: 5}
    for module in ("2", "3"):
        for shelf_no in range(1, 6):
            width_cm = 200.0 if module == "3" and shelf_no == 5 else (
                100.0 if module == "3" else 66.5
            )
            shelf_rows.append((
                module, shelf_no, width_cm, multipliers[shelf_no],
                priority_ranks[shelf_no], int(module == "3" and shelf_no == 5),
            ))
    cursor.executemany("""
        INSERT OR REPLACE INTO chocolate_shelves
        (module, shelf_no, width_cm, multiplier, priority_rank, package_only)
        VALUES (?, ?, ?, ?, ?, ?)
    """, shelf_rows)
    cursor.executemany("""
        INSERT OR REPLACE INTO chocolate_scoring_weights (metric, weight)
        VALUES (?, ?)
    """, DEFAULT_CHOCOLATE_WEIGHTS.items())

    conn.commit()


def clean_str(val):
    if val is None:
        return ""
    return str(val).strip()


def migrate_gondolbasi(conn):
    path = os.path.join(DATA_DIR, "gondolbasi.xlsx")
    if not os.path.exists(path):
        raise FileNotFoundError(f"gondolbasi.xlsx bulunamadı: {path}")

    print("gondolbasi.xlsx dosyası okunuyor...")
    wb = openpyxl.load_workbook(path, data_only=True)

    # 1. Ürün Listesi sheet'ini oku
    ws_ul = wb["Ürün Listesi"]
    ul_rows = list(ws_ul.iter_rows(values_only=True))
    headers_ul = [clean_str(h) for h in ul_rows[0]]
    
    prod_name_idx = headers_ul.index("Ürün Adı")
    cluster_idx = headers_ul.index("Ürün Listesi")
    shelf_idx = headers_ul.index("Raf")
    width_idx = headers_ul.index("Ürün_cm")

    gondol_products = []
    for row in ul_rows[1:]:
        if row[prod_name_idx]:
            gondol_products.append((
                clean_str(row[prod_name_idx]),
                clean_str(row[cluster_idx]),
                row[shelf_idx],
                float(row[width_idx]) if row[width_idx] is not None else 0.0
            ))

    cursor = conn.cursor()
    cursor.executemany("""
    INSERT OR REPLACE INTO gondol_products (product_name, cluster, default_shelf, width_cm)
    VALUES (?, ?, ?, ?)
    """, gondol_products)
    print(f"  ✓ {len(gondol_products)} gondol ürünü eklendi.")

    # 2. 2025 (Satış) sheet'ini oku
    ws_2025 = wb["2025"]
    rows_2025 = list(ws_2025.iter_rows(values_only=True))
    headers_2025 = [clean_str(h) for h in rows_2025[0]]

    ay_idx = headers_2025.index("AY")
    stn_name_idx = headers_2025.index("İstasyon Adı")
    roc_idx = headers_2025.index("Roc Kodu")
    desc_idx = headers_2025.index("Malzeme Açıklaması")
    qty_idx = headers_2025.index("Net Satış Miktarı")

    stations = {}
    sales_agg = {}

    for row in rows_2025[1:]:
        ay = clean_str(row[ay_idx])
        stn_ad = clean_str(row[stn_name_idx])
        roc = row[roc_idx]
        mal_ad = clean_str(row[desc_idx])
        qty = row[qty_idx]

        if not (roc and mal_ad and ay):
            continue

        quarter = QUARTER_MAP.get(ay)
        if not quarter:
            continue

        roc = int(roc)
        qty = float(qty) if qty else 0.0
        
        stations[roc] = stn_ad
        
        key = (roc, quarter, mal_ad)
        sales_agg[key] = sales_agg.get(key, 0.0) + qty

    # İstasyonları ekle
    stations_data = [(roc, name) for roc, name in stations.items()]
    cursor.executemany("""
    INSERT OR REPLACE INTO stations (roc, name)
    VALUES (?, ?)
    """, stations_data)
    print(f"  ✓ {len(stations_data)} istasyon eklendi.")

    # Satışları ekle
    sales_data = [(k[0], k[1], k[2], v) for k, v in sales_agg.items()]
    cursor.executemany("""
    INSERT OR REPLACE INTO sales (roc, quarter, product_name, qty)
    VALUES (?, ?, ?, ?)
    """, sales_data)
    print(f"  ✓ {len(sales_data)} satış kaydı eklendi (çeyreklik olarak toplandı).")

    conn.commit()


def migrate_cikolata(conn):
    path = os.path.join(DATA_DIR, "cikolata.xlsx")
    if not os.path.exists(path):
        raise FileNotFoundError(f"cikolata.xlsx bulunamadı: {path}")

    print("cikolata.xlsx dosyası okunuyor...")
    wb = openpyxl.load_workbook(path, data_only=True)

    # 1. DB Sheet'ini oku (Ürün listesi ve yerleşim tipleri)
    ws_db = wb["DB"]
    db_rows = list(ws_db.iter_rows(values_only=True))
    headers_db = [clean_str(h) for h in db_rows[0]]

    db_id_idx = headers_db.index("Ürün ID")
    db_sku_idx = headers_db.index("SKU")
    db_brand_idx = headers_db.index("Marka")
    db_cat_idx = headers_db.index("Alt Kategori")
    db_width_idx = headers_db.index("Genişlik (cm)")
    db_score_idx = headers_db.index("SKOR")
    db_yerlesim_idx = headers_db.index("Yerleşim")
    db_tip_idx = headers_db.index("Ürün Tipi")

    db_products = {}
    for row in db_rows[1:]:
        u_id = clean_str(row[db_id_idx])
        if not u_id:
            continue
        db_products[u_id] = {
            "urun_id": u_id,
            "short_sku": clean_str(row[db_sku_idx]),
            "brand": clean_str(row[db_brand_idx]),
            "sub_category": clean_str(row[db_cat_idx]),
            "width": float(row[db_width_idx]) if row[db_width_idx] is not None else 0.0,
            "score": float(row[db_score_idx]) if row[db_score_idx] is not None else 0.0,
            "yerlesim": clean_str(row[db_yerlesim_idx]),
            "urun_tipi": clean_str(row[db_tip_idx])
        }

    # 2. Veri_Skor Sheet'ini oku (Metrikler ve full SKU isimleri)
    ws_vs = wb["Veri_Skor"]
    vs_rows = list(ws_vs.iter_rows(values_only=True))
    headers_vs = [clean_str(h) for h in vs_rows[6]]  # Metrik başlıkları 7. satırda

    vs_id_idx = headers_vs.index("Ürün ID")
    vs_sku_idx = headers_vs.index("SKU")  # Full SKU adı (gramajlı)
    vs_satis_idx = headers_vs.index("Satış adeti")
    vs_kar_idx = headers_vs.index("Net Birim Kâr")
    vs_c2_idx = headers_vs.index("C2")
    vs_ciro_idx = headers_vs.index("Ciro Payı")
    vs_gas_idx = headers_vs.index("Nielsen GasSt")
    vs_spm_idx = headers_vs.index("Nielsen SPM")
    vs_deli_idx = headers_vs.index("Deli2go")
    vs_max_facing_idx = headers_vs.index("Max Facing")
    vs_current_facing_idx = headers_vs.index("Mevcut önyüz")
    vs_min_facing_idx = headers_vs.index("Min Facing")
    vs_type_idx = headers_vs.index("Ürün Tipi")

    def get_float(value):
        try:
            return float(value) if value is not None else 0.0
        except (TypeError, ValueError):
            return 0.0

    products_data = []
    metrics_data = []

    for row in vs_rows[7:]:  # 8. satırdan itibaren ürün verileri başlıyor
        u_id = clean_str(row[vs_id_idx])
        if not u_id or u_id == "None":
            continue

        sku_full = clean_str(row[vs_sku_idx])
        db_info = db_products.get(u_id, {
            "urun_id": u_id, "short_sku": sku_full, "brand": "", "sub_category": "Chocolate",
            "width": 5.0, "score": 0.0, "yerlesim": "", "urun_tipi": "Mevcut"
        })

        min_facing = max(1, int(get_float(row[vs_min_facing_idx]) or 1))
        max_facing = max(min_facing, int(get_float(row[vs_max_facing_idx]) or min_facing))
        current_facing = max(0, int(get_float(row[vs_current_facing_idx])))
        product_type = clean_str(row[vs_type_idx]) or db_info["urun_tipi"] or "Mevcut"
        is_package = int(u_id in PACKAGE_PRODUCT_IDS)

        products_data.append((
            u_id,
            sku_full,
            db_info["short_sku"],
            db_info["brand"],
            db_info["sub_category"],
            db_info["width"],
            db_info["score"],
            db_info["yerlesim"],
            product_type,
            min_facing,
            max_facing,
            current_facing,
            is_package,
        ))

        metrics_data.append((
            u_id,
            get_float(row[vs_satis_idx]),
            get_float(row[vs_kar_idx]),
            get_float(row[vs_c2_idx]),
            get_float(row[vs_ciro_idx]),
            get_float(row[vs_gas_idx]),
            get_float(row[vs_spm_idx]),
            get_float(row[vs_deli_idx])
        ))

    cursor = conn.cursor()
    cursor.executemany("""
    INSERT OR REPLACE INTO chocolate_products
    (urun_id, sku, short_sku, brand, sub_category, width_cm, score, yerlesim,
     urun_tipi, min_facing, max_facing, current_facing, is_package)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, products_data)
    print(f"  ✓ {len(products_data)} çikolata ürünü tanımlandı.")

    cursor.executemany("""
    INSERT OR REPLACE INTO chocolate_metrics (urun_id, satis, birim_kar, c2, ciro_pay, niel_gas, niel_spm, deli2go)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, metrics_data)
    print(f"  ✓ {len(metrics_data)} çikolata metrik kaydı eklendi.")

    # 3. İstasyon bazlı çikolata satış verilerini oku
    # Sayfa adı sayısal değerle başlayan (ROC kodu olan) sayfaları filtrele
    station_sales_count = 0
    for sheetname in wb.sheetnames:
        if sheetname and sheetname[0].isdigit() and " " in sheetname:
            try:
                roc = int(sheetname.split(" ")[0])
            except ValueError:
                continue

            ws_st = wb[sheetname]
            rows_st = list(ws_st.iter_rows(values_only=True))
            headers_st = [clean_str(h) for h in rows_st[1]]  # Başlıklar 2. satırda

            st_id_idx = headers_st.index("Ürün ID")
            q1_idx = headers_st.index("Q1")
            q2_idx = headers_st.index("Q2")
            q3_idx = headers_st.index("Q3")
            q4_idx = headers_st.index("Q4")

            st_sales = []
            for row in rows_st[2:]:
                u_id = clean_str(row[st_id_idx])
                if not u_id or u_id == "None" or u_id == "TOPLAM":
                    continue
                
                st_sales.append((
                    roc,
                    u_id,
                    get_float(row[q1_idx]),
                    get_float(row[q2_idx]),
                    get_float(row[q3_idx]),
                    get_float(row[q4_idx])
                ))

            cursor.executemany("""
            INSERT OR REPLACE INTO chocolate_station_sales (roc, urun_id, q1, q2, q3, q4)
            VALUES (?, ?, ?, ?, ?, ?)
            """, st_sales)
            station_sales_count += len(st_sales)

    print(f"  ✓ {station_sales_count} istasyon-çikolata satış kaydı eklendi.")
    conn.commit()


def setup_users(conn):
    cursor = conn.cursor()

    # İstasyon bilgilerini veritabanından çekelim (kullanıcı hesapları oluşturmak için)
    cursor.execute("SELECT roc, name FROM stations")
    stations = cursor.fetchall()

    users_data = []

    # 1. Admin kullanıcısı
    users_data.append((
        ADMIN_USERNAME,
        "admin",
        "Yönetici",
        None
    ))

    # 2. İstasyon kullanıcıları
    def station_username(name):
        s = name.lower().strip()
        for a, b in (
            ("ı", "i"), ("ğ", "g"), ("ü", "u"), ("ş", "s"), ("ö", "o"), ("ç", "c"),
            ("İ", "i"), ("Ğ", "g"), ("Ü", "u"), ("Ş", "s"), ("Ö", "o"), ("Ç", "c"),
        ):
            s = s.replace(a, b)
        import re
        return re.sub(r"[^a-z0-9]", "", s)

    for roc, name in stations:
        uname = station_username(name)
        if not uname:
            continue
        users_data.append((
            uname,
            "station",
            name,
            int(roc)
        ))

    cursor.executemany("""
    INSERT OR REPLACE INTO users (username, role, display_name, roc)
    VALUES (?, ?, ?, ?)
    """, users_data)
    print(f"  ✓ {len(users_data)} kullanıcı hesabı oluşturuldu.")
    conn.commit()


def validate_database(conn):
    integrity = conn.execute("PRAGMA integrity_check").fetchone()[0]
    if integrity != "ok":
        raise RuntimeError(f"SQLite bütünlük kontrolü başarısız: {integrity}")

    foreign_key_errors = conn.execute("PRAGMA foreign_key_check").fetchall()
    if foreign_key_errors:
        raise RuntimeError(f"Foreign key hatası: {foreign_key_errors[:5]}")

    required_counts = {
        "stations": 1,
        "gondol_products": 1,
        "sales": 1,
        "chocolate_products": 1,
        "chocolate_metrics": 1,
        "chocolate_station_sales": 1,
        "chocolate_shelves": 10,
        "chocolate_scoring_weights": 7,
        "users": 1,
    }
    for table, minimum in required_counts.items():
        count = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
        if count < minimum:
            raise RuntimeError(f"{table} tablosu boş veya eksik")

    package_count = conn.execute(
        "SELECT COUNT(*) FROM chocolate_products WHERE is_package = 1"
    ).fetchone()[0]
    if package_count != len(PACKAGE_PRODUCT_IDS):
        raise RuntimeError(
            f"Paket ürün sayısı beklenen {len(PACKAGE_PRODUCT_IDS)}, bulunan {package_count}"
        )

    weight_total = conn.execute(
        "SELECT SUM(weight) FROM chocolate_scoring_weights"
    ).fetchone()[0]
    if abs(weight_total - 1.0) > 1e-9:
        raise RuntimeError(f"Skor ağırlıkları toplamı 1 değil: {weight_total}")


def build_database(target_path=TEMP_DB_PATH):
    if os.path.exists(target_path):
        os.remove(target_path)

    conn = connect_db(target_path)
    try:
        create_tables(conn)
        print("\n[1/3] Gondolbaşı verileri aktarılıyor...")
        migrate_gondolbasi(conn)

        print("\n[2/3] Çikolata verileri aktarılıyor...")
        migrate_cikolata(conn)

        print("\n[3/3] Kullanıcı bilgileri kuruluyor...")
        setup_users(conn)
        validate_database(conn)
        conn.commit()
    except Exception:
        conn.rollback()
        conn.close()
        if os.path.exists(target_path):
            os.remove(target_path)
        raise
    else:
        conn.close()


def main():
    print("=" * 60)
    print("  PLANOGRAM VERİTABANI OLUŞTURMA VE GÖÇ İŞLEMİ")
    print("=" * 60)
    print("-> Şifreler veritabanına kaydedilmeyecek (in-memory kontrolleri kullanılacak).")

    try:
        os.makedirs(DATA_DIR, exist_ok=True)
        build_database()
        os.replace(TEMP_DB_PATH, DB_PATH)
        print("\n" + "=" * 60)
        print("  GÖÇ İŞLEMİ BAŞARIYLA TAMAMLANDI!")
        print("=" * 60)
        return 0
    except Exception as e:
        print(f"\n[HATA] Göç işlemi başarısız oldu: {e}")
        import traceback
        traceback.print_exc()
        if os.path.exists(TEMP_DB_PATH):
            os.remove(TEMP_DB_PATH)
        print("Mevcut planogram.db değiştirilmedi.")
        return 1


if __name__ == "__main__":
    sys.exit(main())
